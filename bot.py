# ============================================
# ЮРИДИЧЕСКИЙ ИИ-АССИСТЕНТ v2.0
# Улучшенное извлечение текста (PDF/DOCX/OCR)
# ============================================

import telebot
import requests
import json
import uuid
import PyPDF2
import pdfplumber
import pytesseract
from PIL import Image
import pdf2image
import docx
import os
import re
from datetime import datetime
from collections import Counter
import warnings
warnings.filterwarnings('ignore')
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ТОКЕНЫ
TELEGRAM_TOKEN = "8448215847:AAEc7o4MSdbMGYXWkb39KNAfKqkaBrGqdZs"
GIGACHAT_AUTH_KEY = "MDE5YzllNWYtNTFhMC03OWFkLTkyYjItYjk4NDA2OTFkZjkxOmZiMWI0NTRjLTMwMDItNDI4NC05Yzk4LTU3NmYzYWRiMWQ1YQ=="

print("✅ Токены загружены")

bot = telebot.TeleBot(TELEGRAM_TOKEN)

SYSTEM_PROMPT = """
Ты - профессиональный юрист-аналитик по B2B договорам.
Проанализируй договор и верни СТРОГО JSON (без markdown):

{
    "стороны": ["Сторона 1", "Сторона 2"],
    "предмет_договора": "краткое описание",
    "сумма_и_валюта": "сумма",
    "сроки_исполнения": "ключевые даты",
    "порядок_оплаты": "условия",
    "штрафные_санкции": [
        {
            "вид_нарушения": "тип",
            "размер_штрафа": "размер",
            "степень_опасности": "критическая/высокая/средняя/низкая"
        }
    ],
    "рисковые_формулировки": [
        {
            "пункт_договора": "цитата",
            "описание_риска": "в чем опасность",
            "категория_риска": "финансовый/временной/юридический/репутационный",
            "уровень_критичности": "критический/высокий/средний/низкий",
            "рекомендация": "что изменить"
        }
    ],
    "ограничения_ответственности": ["пункт1"],
    "условия_расторжения": "описание",
    "подсудность": "где споры",
    "общий_уровень_риска": "критический/высокий/средний/низкий",
    "ключевые_рекомендации": ["рек1", "рек2"],
    "чек_лист_проверки": ["проверить1", "проверить2"]
}

Особое внимание: "но не более...", "за исключением...", "в одностороннем порядке", "без согласования", "неустойка в размере", ограничения ответственности, скрытые платежи.
"""

# Извлечение текста из PDF (3 метода)
def extract_with_pypdf2(file_path):
    try:
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            return "\n".join([page.extract_text() or "" for page in reader.pages])
    except:
        return ""

def extract_with_pdfplumber(file_path):
    try:
        text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        text += " | ".join([str(cell) if cell else "" for cell in row]) + "\n"
        return text
    except:
        return ""

def extract_with_ocr(file_path):
    try:
        images = pdf2image.convert_from_path(file_path, dpi=200, first_page=1, last_page=5)
        text = ""
        for image in images:
            text += pytesseract.image_to_string(image, lang='rus+eng') + "\n"
        return text
    except:
        return ""

def extract_text_from_pdf(file_path):
    # Метод 1: PyPDF2
    text = extract_with_pypdf2(file_path)
    if text and len(text.strip()) > 100:
        print(f"✅ PyPDF2: {len(text)} символов")
        return text
    
    # Метод 2: pdfplumber
    text = extract_with_pdfplumber(file_path)
    if text and len(text.strip()) > 100:
        print(f"✅ pdfplumber: {len(text)} символов")
        return text
    
    # Метод 3: OCR
    print("🔄 Пробую OCR...")
    text = extract_with_ocr(file_path)
    if text and len(text.strip()) > 50:
        print(f"✅ OCR: {len(text)} символов")
        return text
    
    return ""

def extract_text_from_docx(file_path):
    try:
        doc = docx.Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
        for table in doc.tables:
            for row in table.rows:
                text += "\n" + " | ".join([cell.text for cell in row.cells])
        return text
    except:
        return ""

# GigaChat
def get_gigachat_token():
    url = "https://ngw.devices.sberbank.ru:9443/api/v2/oauth"
    headers = {
        'Content-Type': 'application/x-www-form-urlencoded',
        'Accept': 'application/json',
        'RqUID': str(uuid.uuid4()),
        'Authorization': f'Basic {GIGACHAT_AUTH_KEY}'
    }
    try:
        response = requests.post(url, headers=headers, data='scope=GIGACHAT_API_PERS', verify=False, timeout=30)
        if response.status_code == 200:
            return response.json().get('access_token')
        return None
    except:
        return None

def analyze_contract(text):
    token = get_gigachat_token()
    if not token:
        return None
    
    payload = {
        "model": "GigaChat",
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": f"Договор:\n\n{text[:40000]}"}
        ],
        "temperature": 0.2,
        "max_tokens": 4000
    }
    
    headers = {'Content-Type': 'application/json', 'Authorization': f'Bearer {token}'}
    
    try:
        response = requests.post(
            "https://gigachat.devices.sberbank.ru/api/v1/chat/completions",
            headers=headers, json=payload, verify=False, timeout=60
        )
        if response.status_code == 200:
            content = response.json()['choices'][0]['message']['content']
            json_match = re.search(r'\{.*\}', content, re.DOTALL)
            if json_match:
                return json.loads(json_match.group())
        return None
    except:
        return None

# Форматирование
def format_result(analysis):
    if not analysis:
        return "❌ Не удалось проанализировать договор"
    
    risk_emoji = {"критический": "🔴", "высокий": "🟠", "средний": "🟡", "низкий": "🟢"}
    overall = analysis.get('общий_уровень_риска', 'средний')
    emoji = risk_emoji.get(overall, '⚪')
    
    lines = [
        f"{emoji} **АНАЛИЗ ДОГОВОРА**",
        f"**Общий уровень риска: {overall.upper()}**",
        "═" * 30, "",
        "📋 **ОСНОВНАЯ ИНФОРМАЦИЯ:**",
        f"• Стороны: {', '.join(analysis.get('стороны', ['Н/Д']))}",
        f"• Предмет: {analysis.get('предмет_договора', 'Н/Д')}",
        f"• Сумма: {analysis.get('сумма_и_валюта', 'Н/Д')}",
        f"• Сроки: {analysis.get('сроки_исполнения', 'Н/Д')}",
        f"• Оплата: {analysis.get('порядок_оплаты', 'Н/Д')}",
    ]
    
    if analysis.get('подсудность'):
        lines.append(f"• Подсудность: {analysis['подсудность']}")
    
    penalties = analysis.get('штрафные_санкции', [])
    if penalties:
        lines.append(""); lines.append("⚡ **ШТРАФНЫЕ САНКЦИИ:**")
        for p in penalties:
            de = {"критическая": "🔴", "высокая": "🟠", "средняя": "🟡", "низкая": "🟢"}
            lines.append(f"{de.get(p.get('степень_опасности', ''), '⚪')} {p.get('вид_нарушения', '')}: {p.get('размер_штрафа', '')}")
    
    risks = analysis.get('рисковые_формулировки', [])
    if risks:
        lines.append(""); lines.append(f"⚠️ **НАЙДЕНО РИСКОВ: {len(risks)}**")
        for level in ['критический', 'высокий', 'средний', 'низкий']:
            filtered = [r for r in risks if r.get('уровень_критичности') == level]
            if filtered:
                le = {"критический": "🔴 КРИТИЧЕСКИЕ", "высокий": "🟠 ВЫСОКИЕ", "средний": "🟡 СРЕДНИЕ", "низкий": "🟢 НИЗКИЕ"}
                lines.append(""); lines.append(f"**{le.get(level, level.upper())}:**")
                for i, r in enumerate(filtered, 1):
                    lines.append(f"{i}. {r.get('пункт_договора', '')[:150]}")
                    lines.append(f"   ⚠️ {r.get('описание_риска', '')}")
                    if r.get('рекомендация'): lines.append(f"   💡 {r['рекомендация']}")
    
    if analysis.get('условия_расторжения'):
        lines.append(""); lines.append(f"🚪 **Расторжение:** {analysis['условия_расторжения']}")
    
    limitations = analysis.get('ограничения_ответственности', [])
    if limitations:
        lines.append(""); lines.append("🛡️ **Ограничения ответственности:**")
        for l in limitations: lines.append(f"• {l}")
    
    recs = analysis.get('ключевые_рекомендации', [])
    if recs:
        lines.append(""); lines.append("💡 **РЕКОМЕНДАЦИИ:**")
        for i, r in enumerate(recs, 1): lines.append(f"{i}. {r}")
    
    checklist = analysis.get('чек_лист_проверки', [])
    if checklist:
        lines.append(""); lines.append("✅ **ЧЕК-ЛИСТ:**")
        for item in checklist: lines.append(f"☐ {item}")
    
    if risks:
        cats = Counter(r.get('категория_риска', 'прочее') for r in risks)
        scores = {"критический": 3, "высокий": 2, "средний": 1, "низкий": 0}
        total = sum(scores.get(r.get('уровень_критичности', 'средний'), 1) for r in risks)
        danger_index = (total / (len(risks) * 3) * 100) if risks else 0
        lines.append(""); lines.append("📊 **СТАТИСТИКА:**")
        for cat, count in cats.most_common(): lines.append(f"• {cat.title()}: {count}")
        lines.append(f"📈 Индекс опасности: {danger_index:.0f}%")
        if danger_index > 70: lines.append("🔴 Критический уровень")
        elif danger_index > 40: lines.append("🟡 Средний уровень")
        else: lines.append("🟢 Приемлемый уровень")
    
    return "\n".join(lines)

# Обработка документа
def process_document_file(message):
    user_id = message.from_user.id
    file_name = message.document.file_name
    
    print(f"\n📥 Файл: {file_name}")
    
    if not (file_name.lower().endswith('.pdf') or file_name.lower().endswith('.docx')):
        bot.reply_to(message, "❌ Только PDF и DOCX")
        return
    
    status_msg = bot.reply_to(message, "⏳ Загружаю...")
    
    try:
        file_info = bot.get_file(message.document.file_id)
        downloaded = bot.download_file(file_info.file_path)
        
        file_path = f"/tmp/{user_id}_{file_name}"
        with open(file_path, 'wb') as f:
            f.write(downloaded)
        
        bot.edit_message_text("📖 Извлекаю текст...", message.chat.id, status_msg.message_id)
        
        if file_name.endswith('.pdf'):
            text = extract_text_from_pdf(file_path)
        else:
            text = extract_text_from_docx(file_path)
        
        os.remove(file_path)
        
        if not text or len(text.strip()) < 50:
            bot.edit_message_text(
                "❌ **Не удалось извлечь текст**\n\n"
                "Возможные причины:\n"
                "• PDF защищен паролем\n"
                "• Скан низкого качества\n"
                "• Файл поврежден\n\n"
                "💡 Попробуйте другой файл",
                message.chat.id, status_msg.message_id
            )
            return
        
        print(f"✅ Извлечено: {len(text)} символов")
        
        bot.edit_message_text(f"🤖 Анализирую...\n📄 {len(text)} символов\n⏳ ~30 сек", message.chat.id, status_msg.message_id)
        
        analysis = analyze_contract(text)
        
        if not analysis:
            bot.edit_message_text("❌ Ошибка анализа", message.chat.id, status_msg.message_id)
            return
        
        result = format_result(analysis)
        
        if len(result) > 4000:
            parts = [result[i:i+4000] for i in range(0, len(result), 4000)]
            bot.edit_message_text(parts[0], message.chat.id, status_msg.message_id, parse_mode='Markdown')
            for part in parts[1:]:
                bot.send_message(message.chat.id, part, parse_mode='Markdown')
        else:
            bot.edit_message_text(result, message.chat.id, status_msg.message_id, parse_mode='Markdown')
        
    except Exception as e:
        print(f"❌ Ошибка: {e}")
        try:
            bot.edit_message_text(f"❌ Ошибка: {str(e)[:200]}", message.chat.id, status_msg.message_id)
        except:
            bot.reply_to(message, "❌ Ошибка обработки")

# Команды
@bot.message_handler(commands=['start'])
def start(message):
    bot.reply_to(message, """
🏛️ **ЮРИДИЧЕСКИЙ AI-АССИСТЕНТ v2.0**

**Улучшенное распознавание:**
• Обычные PDF (PyPDF2)
• Сложные PDF (pdfplumber)
• Сканы (OCR Tesseract)
• Word документы

**Отправьте файл договора для анализа!**

/help - справка
    """, parse_mode='Markdown')

@bot.message_handler(commands=['help'])
def help_cmd(message):
    bot.reply_to(message, """
📚 **Помощь:**

1. Отправьте PDF или DOCX
2. Бот извлечет текст (3 метода)
3. GigaChat проанализирует
4. Получите разбор рисков

**Совет:** Лучше текстовые PDF, не сканы.

/start - главная
    """, parse_mode='Markdown')

@bot.message_handler(content_types=['document'])
def handle_document(message):
    process_document_file(message)

@bot.message_handler(func=lambda m: True)
def handle_text(message):
    bot.reply_to(message, "📄 Отправьте PDF или DOCX файл договора.\n/start - информация", parse_mode='Markdown')

# Запуск
if __name__ == "__main__":
    print("=" * 50)
    print("✅ БОТ v2.0 ЗАПУЩЕН!")
    print("=" * 50)
    bot.infinity_polling(timeout=60, long_polling_timeout=60)
